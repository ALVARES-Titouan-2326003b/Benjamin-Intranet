"""Synchronisation unidirectionnelle des factures vers un classeur Excel."""

import os
import re
import tempfile
import unicodedata
from contextlib import contextmanager
from copy import copy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from django.utils import timezone
from openpyxl import Workbook, load_workbook


CANONICAL_COLUMNS = {
    "id": "ID Intranet",
    "numero_facture": "N° facture",
    "societe": "Société",
    "affaire": "Affaire",
    "dossier": "Dossier",
    "service": "Pôle",
    "fournisseur": "Fournisseur",
    "montant": "Montant TTC",
    "statut": "Statut",
    "date_facture": "Date facture",
    "echeance": "Échéance",
    "priorite": "Priorité",
    "demandeur": "Demandeur",
    "commentaire": "Commentaire",
    "titre": "Titre",
    "date_soumission": "Date de soumission",
}

COLUMN_ALIASES = {
    "id": {"id", "id facture", "id intranet", "reference intranet"},
    "numero_facture": {"numero facture", "n facture", "no facture", "facture"},
    "societe": {"societe", "societe concernee"},
    "affaire": {"affaire", "affaire concernee"},
    "dossier": {
        "dossier",
        "dossier affaire",
        "dossier concerne",
        "dossier affaire concerne",
    },
    "service": {"pole", "service"},
    "fournisseur": {"fournisseur"},
    "montant": {"montant", "montant ttc", "montant ttc eur", "montant eur"},
    "statut": {"statut", "etat"},
    "date_facture": {"date facture", "date de facture"},
    "echeance": {"echeance", "echeance de paiement", "date echeance"},
    "priorite": {"priorite"},
    "demandeur": {"demandeur", "collaborateur", "ajoute par"},
    "commentaire": {"commentaire", "commentaire compta", "commentaire pour la compta"},
    "titre": {"titre", "objet"},
    "date_soumission": {"date soumission", "date de soumission", "date ajout", "cree le"},
}

REGISTRY_SHEET = "_IntranetSync"


@dataclass(frozen=True)
class ExcelSyncResult:
    workbook_path: Path
    sheet_name: str
    row: int
    created: bool


def _normalize_header(value):
    value = unicodedata.normalize("NFKD", str(value or ""))
    value = "".join(char for char in value if not unicodedata.combining(char))
    value = re.sub(r"[^a-z0-9]+", " ", value.lower()).strip()
    return value


def _invoice_values(invoice):
    demandeur = ""
    if invoice.demandeur_id:
        demandeur = invoice.demandeur.get_full_name() or invoice.demandeur.username
    echeance = invoice.echeance
    if echeance and timezone.is_aware(echeance):
        echeance = timezone.localtime(echeance).replace(tzinfo=None)
    submitted = invoice.date_soumission
    if submitted and timezone.is_aware(submitted):
        submitted = timezone.localtime(submitted).replace(tzinfo=None)
    return {
        "id": invoice.pk,
        "numero_facture": invoice.numero_facture or "",
        "societe": str(invoice.societe) if invoice.societe_id else "",
        "affaire": invoice.affaire or "",
        "dossier": str(invoice.dossier) if invoice.dossier_id else "Divers",
        "service": invoice.get_service_display() if invoice.service else "",
        "fournisseur": str(invoice.fournisseur) if invoice.fournisseur_id else "",
        "montant": invoice.montant,
        "statut": invoice.get_statut_display(),
        "date_facture": invoice.date_facture,
        "echeance": echeance,
        "priorite": invoice.get_priorite_display() if invoice.priorite else "",
        "demandeur": demandeur,
        "commentaire": invoice.commentaire_compta or "",
        "titre": invoice.titre or "",
        "date_soumission": submitted,
    }


@contextmanager
def _workbook_lock(path):
    """Verrou interprocessus placé à côté du classeur partagé."""
    lock_path = path.with_suffix(path.suffix + ".lock")
    lock_path.parent.mkdir(parents=True, exist_ok=True)
    with lock_path.open("a+b") as lock_file:
        if os.name == "nt":
            import msvcrt

            if lock_path.stat().st_size == 0:
                lock_file.write(b"0")
                lock_file.flush()
            lock_file.seek(0)
            msvcrt.locking(lock_file.fileno(), msvcrt.LK_LOCK, 1)
            try:
                yield
            finally:
                lock_file.seek(0)
                msvcrt.locking(lock_file.fileno(), msvcrt.LK_UNLCK, 1)
        else:
            import fcntl

            fcntl.flock(lock_file.fileno(), fcntl.LOCK_EX)
            try:
                yield
            finally:
                fcntl.flock(lock_file.fileno(), fcntl.LOCK_UN)


def _load_or_create_workbook(path, sheet_name, create_if_missing):
    if path.exists():
        return load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm")
    if not create_if_missing:
        raise FileNotFoundError(f"Classeur Excel introuvable : {path}")
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name or "Factures"
    worksheet.append(list(CANONICAL_COLUMNS.values()))
    return workbook


def _worksheet(workbook, sheet_name):
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Feuille Excel introuvable : {sheet_name}")
        return workbook[sheet_name]
    visible_sheets = [sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
    if not visible_sheets:
        raise ValueError("Le classeur ne contient aucune feuille visible.")
    return visible_sheets[0]


def _column_mapping(worksheet, header_row):
    aliases = {
        alias: field
        for field, field_aliases in COLUMN_ALIASES.items()
        for alias in field_aliases
    }
    mapping = {}
    for cell in worksheet[header_row]:
        field = aliases.get(_normalize_header(cell.value))
        if field and field not in mapping:
            mapping[field] = cell.column
    if not mapping:
        raise ValueError("Aucune colonne du fichier Excel ne correspond aux champs de facture.")
    if "id" not in mapping and "numero_facture" not in mapping:
        raise ValueError("Le fichier Excel doit contenir une colonne ID ou N° facture.")
    return mapping


def _registry(workbook):
    if REGISTRY_SHEET in workbook.sheetnames:
        worksheet = workbook[REGISTRY_SHEET]
    else:
        worksheet = workbook.create_sheet(REGISTRY_SHEET)
        worksheet.append(["ID Intranet", "Feuille", "Ligne", "Synchronisé le"])
    worksheet.sheet_state = "hidden"
    return worksheet


def _registered_row(registry, invoice_id, sheet_name):
    for row in registry.iter_rows(min_row=2, values_only=False):
        if str(row[0].value or "") == str(invoice_id) and row[1].value == sheet_name:
            return row[0].row, int(row[2].value)
    return None, None


def _find_existing_row(worksheet, mapping, values, header_row):
    identifiers = []
    if "id" in mapping:
        identifiers.append((mapping["id"], str(values["id"])))
    if values["numero_facture"] and "numero_facture" in mapping:
        identifiers.append((mapping["numero_facture"], str(values["numero_facture"])))
    for row_number in range(header_row + 1, worksheet.max_row + 1):
        for column, expected in identifiers:
            if str(worksheet.cell(row_number, column).value or "").strip() == expected.strip():
                return row_number
    return None


def _copy_row_style(worksheet, source_row, target_row, columns):
    if source_row <= 0:
        return
    for column in columns:
        source = worksheet.cell(source_row, column)
        target = worksheet.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        target.number_format = source.number_format


def _save_atomically(workbook, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    suffix = path.suffix or ".xlsx"
    handle, temporary_name = tempfile.mkstemp(prefix=f".{path.stem}-", suffix=suffix, dir=path.parent)
    os.close(handle)
    temporary_path = Path(temporary_name)
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, path)
    finally:
        temporary_path.unlink(missing_ok=True)


def sync_invoice_to_workbook(
    invoice,
    workbook_path,
    *,
    sheet_name="",
    header_row=1,
    create_if_missing=False,
):
    """Ajoute ou met à jour une facture dans le classeur configuré."""
    path = Path(workbook_path).expanduser().resolve()
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("La synchronisation accepte uniquement les fichiers .xlsx ou .xlsm.")
    if header_row < 1:
        raise ValueError("La ligne d'en-tête Excel doit être supérieure ou égale à 1.")

    with _workbook_lock(path):
        workbook = _load_or_create_workbook(path, sheet_name, create_if_missing)
        worksheet = _worksheet(workbook, sheet_name)
        mapping = _column_mapping(worksheet, header_row)
        registry = _registry(workbook)
        values = _invoice_values(invoice)

        registry_row, target_row = _registered_row(registry, invoice.pk, worksheet.title)
        if target_row and target_row <= header_row:
            target_row = None
        if not target_row:
            target_row = _find_existing_row(worksheet, mapping, values, header_row)
        created = target_row is None
        if created:
            target_row = max(worksheet.max_row + 1, header_row + 1)
            _copy_row_style(worksheet, target_row - 1, target_row, mapping.values())

        for field, column in mapping.items():
            worksheet.cell(target_row, column).value = values[field]

        synchronized_at = datetime.now().isoformat(timespec="seconds")
        if registry_row:
            registry.cell(registry_row, 3).value = target_row
            registry.cell(registry_row, 4).value = synchronized_at
        else:
            registry.append([invoice.pk, worksheet.title, target_row, synchronized_at])

        _save_atomically(workbook, path)
        return ExcelSyncResult(path, worksheet.title, target_row, created)
