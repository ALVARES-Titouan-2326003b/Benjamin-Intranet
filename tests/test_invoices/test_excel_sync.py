from datetime import date
from io import StringIO

import pytest
from django.contrib.auth.models import User
from django.core.management import call_command
from openpyxl import Workbook, load_workbook

from invoices.models import ActeurExterne, Client, Facture, Fournisseur
from invoices.services.excel_sync import REGISTRY_SHEET, sync_invoice_to_workbook
from invoices.views import synchronize_configured_invoice_workbook


@pytest.fixture
def excel_invoice(db, societe):
    user = User.objects.create_user(
        username="excel-owner",
        first_name="Alice",
        last_name="Martin",
    )
    supplier_actor = ActeurExterne.objects.create(id="EXCEL-SUPPLIER")
    supplier = Fournisseur.objects.create(id=supplier_actor, nom="Entreprise Excel")
    client_actor = ActeurExterne.objects.create(id="EXCEL-CLIENT")
    client = Client.objects.create(id=client_actor)
    return Facture.objects.create(
        id="FAC-EXCEL-001",
        numero_facture="F-2026-001",
        societe=societe,
        fournisseur=supplier,
        client=client,
        montant=1250.5,
        statut="received",
        service="administratif",
        date_facture=date(2026, 8, 12),
        priorite="urgent",
        commentaire_compta="À contrôler",
        titre="Honoraires",
        demandeur=user,
        created_by=user,
    )


def _existing_workbook(path):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Suivi"
    worksheet.append(
        [
            "N° facture",
            "Société",
            "Dossier / affaire concerné",
            "Pôle",
            "Fournisseur",
            "Montant TTC (€)",
            "Statut",
            "Date de facture",
            "Priorité",
            "Demandeur",
            "Commentaire pour la compta",
            "Titre",
            "Colonne libre",
        ]
    )
    worksheet.append(["ANCIENNE", "", "", "", "", 10, "", "", "", "", "", "", "préservée"])
    workbook.save(path)


@pytest.mark.django_db
def test_sync_adds_then_updates_the_same_invoice_row(tmp_path, excel_invoice):
    workbook_path = tmp_path / "suivi.xlsx"
    _existing_workbook(workbook_path)

    first = sync_invoice_to_workbook(
        excel_invoice,
        workbook_path,
        sheet_name="Suivi",
    )

    assert first.created is True
    assert first.row == 3
    workbook = load_workbook(workbook_path)
    worksheet = workbook["Suivi"]
    assert worksheet.cell(3, 1).value == "F-2026-001"
    assert worksheet.cell(3, 2).value == "Benjamin Immobilier"
    assert worksheet.cell(3, 5).value == "Entreprise Excel"
    assert worksheet.cell(3, 6).value == 1250.5
    assert worksheet.cell(3, 13).value is None
    assert workbook[REGISTRY_SHEET].sheet_state == "hidden"

    excel_invoice.numero_facture = "F-2026-001-CORRIGE"
    excel_invoice.montant = 1450
    excel_invoice.save(update_fields=["numero_facture", "montant"])
    second = sync_invoice_to_workbook(
        excel_invoice,
        workbook_path,
        sheet_name="Suivi",
    )

    assert second.created is False
    assert second.row == 3
    workbook = load_workbook(workbook_path)
    worksheet = workbook["Suivi"]
    assert worksheet.max_row == 3
    assert worksheet.cell(3, 1).value == "F-2026-001-CORRIGE"
    assert worksheet.cell(3, 6).value == 1450
    assert worksheet.cell(2, 13).value == "préservée"


@pytest.mark.django_db
def test_sync_can_create_a_new_workbook_when_explicitly_allowed(tmp_path, excel_invoice):
    workbook_path = tmp_path / "nouveau.xlsx"

    result = sync_invoice_to_workbook(
        excel_invoice,
        workbook_path,
        create_if_missing=True,
    )

    assert result.created is True
    workbook = load_workbook(workbook_path)
    assert workbook["Factures"].cell(2, 1).value == excel_invoice.pk
    assert workbook["Factures"].cell(2, 2).value == "F-2026-001"


@pytest.mark.django_db
def test_configured_sync_records_success_and_failure(settings, tmp_path, excel_invoice):
    settings.INVOICE_EXCEL_SYNC_PATH = str(tmp_path / "suivi.xlsx")
    settings.INVOICE_EXCEL_SYNC_SHEET = "Factures"
    settings.INVOICE_EXCEL_SYNC_HEADER_ROW = 1
    settings.INVOICE_EXCEL_SYNC_CREATE_IF_MISSING = True

    assert synchronize_configured_invoice_workbook(excel_invoice, excel_invoice.created_by) is True
    success = excel_invoice.historique.get(action="excel_sync_success")
    assert "ligne 2" in success.details

    settings.INVOICE_EXCEL_SYNC_SHEET = "Feuille absente"
    assert synchronize_configured_invoice_workbook(excel_invoice, excel_invoice.created_by) is False
    error = excel_invoice.historique.get(action="excel_sync_error")
    assert "Feuille Excel introuvable" in error.details


@pytest.mark.django_db
def test_unconfigured_sync_is_a_noop(settings, excel_invoice):
    settings.INVOICE_EXCEL_SYNC_PATH = ""

    assert synchronize_configured_invoice_workbook(excel_invoice) is None
    assert not excel_invoice.historique.filter(action__startswith="excel_sync").exists()


@pytest.mark.django_db
def test_management_command_resynchronizes_existing_invoices(settings, tmp_path, excel_invoice):
    settings.INVOICE_EXCEL_SYNC_PATH = str(tmp_path / "reprise.xlsx")
    settings.INVOICE_EXCEL_SYNC_SHEET = "Factures"
    settings.INVOICE_EXCEL_SYNC_HEADER_ROW = 1
    settings.INVOICE_EXCEL_SYNC_CREATE_IF_MISSING = True
    stdout = StringIO()

    call_command("sync_invoices_excel", stdout=stdout)

    assert "1/1 réussie" in stdout.getvalue()
    workbook = load_workbook(settings.INVOICE_EXCEL_SYNC_PATH)
    assert workbook["Factures"].cell(2, 1).value == excel_invoice.pk
